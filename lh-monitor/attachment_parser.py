"""Extract searchable text from LH attachments; unknown or unreadable files need review."""

from __future__ import annotations

import io
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from parser import DIRECT, clean

MAX_UNCOMPRESSED = 120_000_000
SUPPORTED = {".pdf", ".hwp", ".hwpx", ".xls", ".xlsx", ".zip"}
PDF_LOCK = threading.Lock()


def should_inspect(name: str) -> bool:
    return Path(name).suffix.lower() in SUPPORTED


def extract_text(name: str, data: bytes, depth: int = 0) -> str:
    if len(data) > MAX_UNCOMPRESSED or depth > 2:
        raise ValueError("첨부파일 크기 또는 압축 깊이 제한 초과")
    suffix = Path(name).suffix.lower()
    if data.startswith(b"<!DOCTYPE html") or data.startswith(b"<html"):
        raise ValueError("첨부파일 대신 HTML 오류 화면 수신")
    if suffix == ".pdf":
        import pymupdf
        try:
            with PDF_LOCK:
                with pymupdf.open(stream=data, filetype="pdf") as document:
                    if document.page_count > 300:
                        raise ValueError("PDF 300페이지 초과")
                    result = " ".join(page.get_text() or "" for page in document)
        except Exception as exc:
            raise ValueError(f"PDF 해석 실패: {exc}") from exc
    elif suffix in (".hwpx", ".zip"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED:
                raise ValueError("압축 해제 크기 제한 초과")
            chunks = []
            for item in archive.infolist():
                if item.is_dir():
                    continue
                if suffix == ".hwpx" and item.filename.startswith("Contents/") and item.filename.endswith(".xml"):
                    root = ElementTree.fromstring(archive.read(item))
                    chunks.extend(t for t in root.itertext() if t.strip())
                elif suffix == ".zip" and Path(item.filename).suffix.lower() in (".pdf", ".hwp", ".hwpx", ".xls", ".xlsx", ".zip"):
                    chunks.append(extract_text(item.filename, archive.read(item), depth + 1))
            result = " ".join(chunks)
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        chunks = []
        try:
            for sheet in book:
                for row in sheet.iter_rows(values_only=True):
                    chunks.extend(str(value) for value in row if value is not None)
        finally:
            book.close()
        result = " ".join(chunks)
    elif suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        result = " ".join(str(sheet.cell_value(row, col)) for sheet in book.sheets() for row in range(sheet.nrows) for col in range(sheet.ncols))
    elif suffix == ".hwp":
        executable = shutil.which("hwp5txt")
        if not executable:
            sibling = Path(sys.executable).with_name("hwp5txt.exe" if sys.platform == "win32" else "hwp5txt")
            executable = str(sibling) if sibling.exists() else None
        if not executable:
            raise ValueError("HWP 해석기 hwp5txt 없음")
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "notice.hwp"
            path.write_bytes(data)
            try:
                command = subprocess.run([executable, str(path)], capture_output=True, timeout=30, check=True)
            except subprocess.TimeoutExpired as exc:
                raise ValueError("HWP 해석 30초 제한 초과") from exc
            except subprocess.CalledProcessError as exc:
                raise ValueError(f"HWP 해석기 실패 (exit {exc.returncode})") from exc
            result = command.stdout.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"지원하지 않는 첨부 형식: {suffix}")
    result = clean(result)
    if not result:
        raise ValueError(f"{suffix}에서 추출된 텍스트 없음 (스캔 문서 가능)")
    return result
